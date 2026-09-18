"""周研判报告：从看板数据生成 HTML（自包含、内联 SVG）与 DOCX（python-docx）。

定位：区域级资源配置辅助 —— 汇总本周案件概览、Top 热点网格、疑似系列案组、
周突增预警与处置建议；全部基于合成数据，不含任何个体画像或预测。
"""

from __future__ import annotations

import json
from datetime import date, timedelta
from pathlib import Path

DISCLAIMER = "本报告基于合成数据（虚构城市「滨江市」）自动生成，仅用于技术演示，不构成任何现实研判依据。"


def _load_payloads(data_dir: Path) -> dict:
    payload_dir = data_dir / "web" / "data"
    out: dict = {}
    for name in ("meta", "trend", "frames", "alerts", "groups"):
        path = payload_dir / f"{name}.json"
        out[name] = json.loads(path.read_text(encoding="utf-8")) if path.exists() else None
    return out


def _top_hot_cells(frames_doc: dict, week_index: int, top_n: int = 8) -> list[dict]:
    frame = frames_doc["frames"][max(0, min(week_index, len(frames_doc["frames"]) - 1))]
    nx, cell = frames_doc["nx"], frames_doc["cell_m"]
    cells = sorted(frame["cells"], key=lambda pair: pair[1], reverse=True)[:top_n]
    out = []
    for flat, count in cells:
        row, col = divmod(flat, nx)
        out.append(
            {
                "grid": f"r{row:02d}c{col:02d}",
                "x": round((col + 0.5) * cell, 1),
                "y": round((row + 0.5) * cell, 1),
                "n": int(count),
            }
        )
    return out


def _week_index_for(frames_doc: dict, meta: dict, weeks_back: int) -> int:
    return max(0, len(frames_doc["frames"]) - 1 - weeks_back)


def _week_label(meta: dict, frames_doc: dict, index: int) -> str:
    frame = frames_doc["frames"][index]
    start = date.fromisoformat(meta["start_date"])
    d0 = start + timedelta(days=frame["from"])
    d1 = start + timedelta(days=frame["to"])
    return f"{d0} ~ {d1}（滚动 7 日窗口）"


def collect_weekly(data_dir: Path | str, *, weeks_back: int = 0) -> dict:
    data_dir = Path(data_dir)
    payloads = _load_payloads(data_dir)
    meta, trend, frames, alerts_doc, groups_doc = (
        payloads["meta"], payloads["trend"], payloads["frames"], payloads["alerts"], payloads["groups"]
    )
    if not (meta and frames and trend):
        raise FileNotFoundError("看板数据未构建：请先执行 `policesight web build --data-dir <目录>`")

    week_index = _week_index_for(frames, meta, weeks_back)
    week_total = trend["total"][week_index] if week_index < len(trend["total"]) else 0
    prev_total = trend["total"][week_index - 1] if week_index > 0 else week_total
    change = (week_total - prev_total) / prev_total if prev_total else 0.0
    type_breakdown = sorted(
        ((str(t), series[week_index] if week_index < len(series) else 0) for t, series in trend["by_type"].items()),
        key=lambda pair: pair[1],
        reverse=True,
    )
    week_alerts = [
        alert for alert in (alerts_doc["alerts"] if alerts_doc else []) if alert["week_index"] == week_index
    ]
    hot_cells = _top_hot_cells(frames, week_index)
    groups = (groups_doc["top"] if groups_doc else [])[:10]
    suggestions = []
    seen = set()
    for alert in week_alerts:
        key = (alert["district"], alert["type"])
        if key in seen:
            continue
        seen.add(key)
        suggestions.append(alert["advice"])
        if len(suggestions) >= 6:
            break
    if not suggestions:
        suggestions.append("本周未触发突增预警 —— 建议维持常规巡防强度，关注高发时段（见图）。")

    return {
        "city": meta["city"],
        "week_index": week_index,
        "week_label": _week_label(meta, frames, week_index),
        "week_total": week_total,
        "prev_total": prev_total,
        "change": change,
        "type_breakdown": type_breakdown,
        "hot_cells": hot_cells,
        "alerts": week_alerts,
        "alert_acceptance": alerts_doc["acceptance"] if alerts_doc else {},
        "groups": groups,
        "group_evaluation": groups_doc["evaluation"] if groups_doc else {},
        "weekly_series": trend["total"],
        "weeks": trend["weeks"],
        "suggestions": suggestions,
        "disclaimer": DISCLAIMER,
    }


def _svg_bars(series: list[int], highlight: int, width: int = 640, height: int = 120) -> str:
    if not series:
        return ""
    max_v = max(series) or 1
    bar_w = width / len(series)
    parts = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" role="img">']
    for index, value in enumerate(series):
        h = (value / max_v) * (height - 18)
        x = index * bar_w
        color = "#1d4ed8" if index == highlight else "#93c5fd"
        parts.append(f'<rect x="{x + 1:.1f}" y="{height - h:.1f}" width="{bar_w - 2:.1f}" height="{h:.1f}" fill="{color}"/>')
    parts.append(f'<text x="4" y="12" font-size="10" fill="#6b7280">周案件量（高亮=本周，峰值 {max_v}）</text>')
    parts.append("</svg>")
    return "".join(parts)


def render_html(data: dict) -> str:
    rows_types = "".join(f"<tr><td>{name}</td><td>{count}</td></tr>" for name, count in data["type_breakdown"])
    rows_hot = "".join(
        f"<tr><td>{cell['grid']}</td><td>{cell['x']}</td><td>{cell['y']}</td><td>{cell['n']}</td></tr>"
        for cell in data["hot_cells"]
    )
    rows_alerts = "".join(
        f"<tr><td class='lv-{alert['level']}'>{alert['level']}</td><td>{alert['district']}"
        f"{(' ' + alert['grid']) if alert.get('grid') else ''}</td><td>{alert['type']}</td>"
        f"<td>{alert['current']} / {alert['baseline']} · {alert['factor']}×</td></tr>"
        for alert in data["alerts"][:12]
    ) or "<tr><td colspan='4'>本周无预警</td></tr>"
    rows_groups = "".join(
        f"<tr><td>{group['id']}</td><td>{group['n']}</td><td>{'、'.join(group['methods'])}</td>"
        f"<td>{group['suspect_score']}</td></tr>"
        for group in data["groups"][:10]
    ) or "<tr><td colspan='4'>无可展示分组</td></tr>"
    suggestions = "".join(f"<li>{text}</li>" for text in data["suggestions"])
    change_text = f"{data['change'] * 100:+.1f}%"
    ev = data.get("group_evaluation") or {}
    acceptance = data.get("alert_acceptance") or {}

    return f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><title>PoliceSight 周研判报告 · {data['week_label']}</title>
<style>
 body{{font:14px/1.65 "Microsoft YaHei",system-ui,sans-serif;color:#111827;max-width:860px;margin:24px auto;padding:0 18px}}
 h1{{font-size:21px}} h2{{font-size:16px;margin-top:26px;border-left:4px solid #1d4ed8;padding-left:9px}}
 table{{border-collapse:collapse;width:100%;margin:8px 0}} td,th{{border:1px solid #e5e7eb;padding:5px 9px;text-align:left;font-size:13px}}
 th{{background:#f8fafc}} .lv-red{{color:#dc2626;font-weight:700}} .lv-orange{{color:#f59e0b;font-weight:700}} .lv-yellow{{color:#ca8a04;font-weight:700}}
 .muted{{color:#6b7280;font-size:12.5px}} .kpi{{display:flex;gap:26px;margin:10px 0}} .kpi b{{font-size:20px}}
 footer{{margin-top:30px;border-top:1px solid #e5e7eb;padding-top:10px}}
</style></head><body>
<h1>PoliceSight · 周研判报告</h1>
<p class="muted">{data['city']} ｜ {data['week_label']} ｜ 自动生成 · {DISCLAIMER}</p>
<div class="kpi"><div>本周案件<br><b>{data['week_total']}</b></div>
<div>环比<br><b>{change_text}</b></div>
<div>预警条数<br><b>{len(data['alerts'])}</b></div>
<div>疑组（Top 展示）<br><b>{len(data['groups'])}</b></div></div>
{_svg_bars(data['weekly_series'], data['week_index'])}

<h2>一、按案件类型</h2><table><tr><th>类型</th><th>本周</th></tr>{rows_types}</table>

<h2>二、Top 热点网格（滚动 7 日 · 250m）</h2>
<table><tr><th>网格</th><th>x(m)</th><th>y(m)</th><th>案件数</th></tr>{rows_hot}</table>

<h2>三、周突增预警（按级别）</h2>
<table><tr><th>级别</th><th>区域/网格</th><th>类型</th><th>本周/基线 · 倍数</th></tr>{rows_alerts}</table>
<p class="muted">真值验收：突增场景命中 {acceptance.get('hit', '—')}/{acceptance.get('anomalies', '—')} ·
热点区对齐 {acceptance.get('hotspot_aligned', '—')} · 无关预警 {acceptance.get('false_alerts', '—')}</p>

<h2>四、疑似系列案组 Top10</h2>
<table><tr><th>组</th><th>案件数</th><th>手法</th><th>可疑度</th></tr>{rows_groups}</table>
<p class="muted">对照真值（预设团伙）：纯度 {ev.get('purity', '—')} · 召回 {ev.get('recall', '—')}（规模效应详见 docs/benchmark-link-alert.md）</p>

<h2>五、处置建议</h2><ul>{suggestions}</ul>
<footer class="muted">由 PoliceSight v0.1.0 生成 ｜ 合成数据 · 离线可复现 ｜ 本报告不涉及个体画像与预测</footer>
</body></html>"""


def render_docx(data: dict, out_path: Path) -> Path:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading("PoliceSight · 周研判报告", level=0)
    doc.add_paragraph(f"{data['city']} ｜ {data['week_label']} ｜ {DISCLAIMER}").runs[0].font.size = Pt(10)

    doc.add_heading("一、概览", level=1)
    doc.add_paragraph(
        f"本周案件 {data['week_total']} 起（环比 {data['change'] * 100:+.1f}%）；"
        f"预警 {len(data['alerts'])} 条；疑组 Top 展示 {len(data['groups'])} 组。"
    )
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "案件类型", "本周案件数"
    for name, count in data["type_breakdown"]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = name, str(count)

    doc.add_heading("二、Top 热点网格（滚动 7 日 · 250m）", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, ("网格", "x(m)", "y(m)", "案件数"), strict=True):
        cell.text = text
    for cell_item in data["hot_cells"]:
        cells = table.add_row().cells
        cells[0].text = cell_item["grid"]
        cells[1].text, cells[2].text, cells[3].text = str(cell_item["x"]), str(cell_item["y"]), str(cell_item["n"])

    doc.add_heading("三、周突增预警", level=1)
    if data["alerts"]:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        for cell, text in zip(table.rows[0].cells, ("级别", "区域/网格", "类型", "本周/基线 · 倍数"), strict=True):
            cell.text = text
        for alert in data["alerts"][:12]:
            cells = table.add_row().cells
            cells[0].text = alert["level"]
            cells[1].text = alert["district"] + (f" {alert['grid']}" if alert.get("grid") else "")
            cells[2].text = alert["type"]
            cells[3].text = f"{alert['current']} / {alert['baseline']} · {alert['factor']}×"
    else:
        doc.add_paragraph("本周无预警。")

    doc.add_heading("四、疑似系列案组 Top10", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, ("组", "案件数", "手法", "可疑度"), strict=True):
        cell.text = text
    for group in data["groups"][:10]:
        cells = table.add_row().cells
        cells[0].text = group["id"]
        cells[1].text = str(group["n"])
        cells[2].text = "、".join(group["methods"])
        cells[3].text = str(group["suspect_score"])

    doc.add_heading("五、处置建议", level=1)
    for text in data["suggestions"]:
        doc.add_paragraph(text, style="List Bullet")

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def build_weekly(data_dir: Path | str, *, weeks_back: int = 0, out_dir: Path | str | None = None) -> dict:
    data_dir = Path(data_dir)
    out = Path(out_dir) if out_dir else data_dir / "reports"
    out.mkdir(parents=True, exist_ok=True)
    weekly = collect_weekly(data_dir, weeks_back=weeks_back)
    stem = f"weekly_report_w{weekly['week_index'] + 1:02d}"
    html_path = out / f"{stem}.html"
    html_path.write_text(render_html(weekly), encoding="utf-8")
    docx_path = render_docx(weekly, out / f"{stem}.docx")
    return {
        "html": str(html_path),
        "docx": str(docx_path),
        "week_index": weekly["week_index"],
        "week_label": weekly["week_label"],
        "cases": weekly["week_total"],
        "alerts": len(weekly["alerts"]),
    }
